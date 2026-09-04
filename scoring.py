"""
Complexity scoring and instance ranking.

Uses prod_transactions CSV as the base (covers all ~11K instances),
and optionally enriches with HI audit xlsx data when available.
Produces a ranked list of the top N most complex instances.
"""
import csv
import json
import os
import re
import xml.etree.ElementTree as ET
from typing import Dict, List, Optional, Tuple

from rich.console import Console

import config

console = Console()


# ── Audit file parsers (reused from V3 methodology) ─────────────────────────

def _parse_row_count(s: str) -> Optional[int]:
    """Extract rowCount from audit Results field."""
    if not s:
        return None
    m = re.search(r'"rowCount"\s*:\s*"?(\d+)"?', s)
    return int(m.group(1)) if m else None


def _parse_flow_count(s: str) -> Optional[int]:
    """Extract High+Medium flow count from audit Results field."""
    if not s:
        return None
    try:
        d = json.loads(s.replace("*** Script: ", "").strip())
        return int(d.get("High", 0)) + int(d.get("Medium", 0))
    except Exception:
        return None


def _load_audit_xlsx(filepath: str, parser_fn) -> Dict[str, int]:
    """Load an audit xlsx file, applying parser_fn to the Results column."""
    try:
        import openpyxl
    except ImportError:
        console.print("[yellow]openpyxl not installed, skipping audit file[/yellow]")
        return {}

    if not os.path.exists(filepath):
        return {}

    data = {}
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        sheet = wb["Export"] if "Export" in wb.sheetnames else wb.active
        rows = list(sheet.iter_rows(values_only=True))
        wb.close()

        for row in rows[1:]:
            if len(row) < 5:
                continue
            inst, res, st = row[2], row[3], row[4]
            if not inst or (st and str(st) != "Completed"):
                continue
            inst = str(inst).strip().lower()
            v = parser_fn(str(res) if res else "")
            if v is not None and isinstance(v, (int, float)):
                data[inst] = max(data.get(inst, 0), v)
    except Exception as e:
        console.print(f"[yellow]Warning: Could not load {filepath}: {e}[/yellow]")

    return data


def _load_table_stats(filepath: str) -> Tuple[Dict[str, Dict[str, int]], List[str]]:
    """Load All_audit_table_statistics.xlsx → {instance: {table: row_count}}."""
    try:
        import openpyxl
    except ImportError:
        return {}, []

    if not os.path.exists(filepath):
        return {}, []

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        sheet_name = "Audit Results" if "Audit Results" in wb.sheetnames else wb.sheetnames[0]
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        wb.close()

        hdrs = [str(h).strip() if h else "" for h in rows[0]]
        tables = hdrs[1:]
        data = {}
        for row in rows[1:]:
            if not row[0]:
                continue
            inst = str(row[0]).strip().lower()
            counts = {}
            for j, t in enumerate(tables, 1):
                if j < len(row) and row[j] is not None:
                    try:
                        v = int(row[j])
                        counts[t] = v if v >= 0 else None
                    except (ValueError, TypeError):
                        counts[t] = None
                else:
                    counts[t] = None
            data[inst] = counts
        return data, tables
    except Exception as e:
        console.print(f"[yellow]Warning: Could not load table stats: {e}[/yellow]")
        return {}, []


# ── CMDB service XML loader ──────────────────────────────────────────────

def load_company_map(filepath: str = None) -> Dict[str, str]:
    """
    Load cmdb_ci_service XML to build {instance_name: company_name} map.
    Service names are formatted as 'SNC Instance - <instance>'.
    """
    fp = filepath or config.CMDB_SERVICE_XML
    if not os.path.exists(fp):
        console.print(f"  [yellow]CMDB XML not found: {fp}[/yellow]")
        return {}

    mapping = {}
    try:
        tree = ET.parse(fp)
        root = tree.getroot()
        for svc in root.findall("cmdb_ci_service"):
            name = svc.findtext("name", "")
            company_el = svc.find("company")
            company = company_el.get("display_value", "") if company_el is not None else ""
            if not name or not company:
                continue
            inst = name.replace("SNC Instance - ", "").strip().lower()
            if inst:
                mapping[inst] = company
        console.print(f"  Loaded [cyan]{len(mapping):,}[/cyan] instance-to-company mappings from CMDB XML")
    except Exception as e:
        console.print(f"  [yellow]Warning: Could not parse CMDB XML: {e}[/yellow]")

    # Supplement with XLSX data for any missing instances
    xlsx_map = _load_company_xlsx()
    added = 0
    for inst, company in xlsx_map.items():
        if inst not in mapping:
            mapping[inst] = company
            added += 1
    if added:
        console.print(f"  Added [cyan]{added}[/cyan] more mappings from CMDB XLSX (total: {len(mapping):,})")

    # Supplement with Prod Customers xlsx for any remaining unmapped instances
    prod_cust = load_prod_customers()
    added2 = 0
    for inst, info in prod_cust.items():
        if inst not in mapping and info.get("company"):
            mapping[inst] = info["company"]
            added2 += 1
    if added2:
        console.print(f"  Added [cyan]{added2:,}[/cyan] more mappings from Prod Customers xlsx (total: {len(mapping):,})")

    return mapping


def _load_company_xlsx(filepath: str = None) -> Dict[str, str]:
    """
    Load cmdb_ci_service XLSX to build {instance_name: company_name} map.
    Column A = Company, Column I = Suffix (instance name).
    Falls back to Column B 'SNC Instance - <name>' if I is empty.
    """
    fp = filepath or getattr(config, "CMDB_SERVICE_XLSX", "")
    if not fp or not os.path.exists(fp):
        return {}

    mapping = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=False):
            vals = {c.column_letter: c.value for c in row if c.value is not None}
            company = vals.get("A", "")
            inst = vals.get("I", "")
            if not inst and vals.get("B", ""):
                inst = str(vals["B"]).replace("SNC Instance - ", "").strip()
            if company and inst:
                mapping[str(inst).strip().lower()] = str(company).strip()
        wb.close()
        console.print(f"  Loaded [cyan]{len(mapping):,}[/cyan] instance-to-company mappings from CMDB XLSX")
    except Exception as e:
        console.print(f"  [yellow]Warning: Could not parse CMDB XLSX: {e}[/yellow]")

    return mapping


# ── Prod Customers XLSX loader ───────────────────────────────────────────────

def load_prod_customers(filepath: str = None) -> Dict[str, dict]:
    """
    Load the Prod Customers List xlsx.
    Returns {instance_name: {company, activity_level, ...}} for all 13K+ prod instances.
    """
    fp = filepath or getattr(config, "PROD_CUSTOMERS_XLSX", "")
    if not fp or not os.path.exists(fp):
        console.print(f"  [yellow]Prod Customers xlsx not found: {fp}[/yellow]")
        return {}

    data = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            suffix = row[9]  # Column J = Suffix (calc) = instance name
            if not suffix:
                continue
            inst = str(suffix).strip().lower()
            data[inst] = {
                "company": str(row[0] or "").strip(),
                "activity_level": str(row[3] or "").strip(),
                "operational_status": str(row[4] or "").strip(),
                "used_for": str(row[6] or "").strip(),
                "type_purpose": str(row[7] or "").strip(),
                "branch": str(row[13] or "").strip(),
            }
        wb.close()
        console.print(f"  Loaded [cyan]{len(data):,}[/cyan] instances from Prod Customers xlsx")
    except Exception as e:
        console.print(f"  [yellow]Warning: Could not parse Prod Customers xlsx: {e}[/yellow]")

    return data


# ── CSV loader ───────────────────────────────────────────────────────────────

def load_prod_transactions(filepath: str = None) -> Dict[str, dict]:
    """
    Load prod_transactions CSV.
    Returns {instance_name: {txn_90d, db_gb, primary_gb, capacity_tier, release, ...}}
    """
    fp = filepath or config.PROD_TRANSACTIONS_CSV
    if not os.path.exists(fp):
        console.print(f"[red]Error: {fp} not found[/red]")
        return {}

    data = {}
    with open(fp, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            inst = r.get("instance", "").strip().lower()
            if not inst:
                continue
            try:
                txn = float(r.get("90_day_average") or 0)
            except (ValueError, TypeError):
                txn = 0
            try:
                db = float(r.get("combined_size") or 0)
            except (ValueError, TypeError):
                db = 0
            try:
                primary = float(r.get("primary_table_size") or 0)
            except (ValueError, TypeError):
                primary = 0

            cap = r.get("u_disco_capacity_size_name", "").strip().lower()
            if cap == "null":
                cap = ""
            data[inst] = {
                "txn_90d": txn,
                "db_gb": db,
                "primary_table_gb": primary,
                "capacity_tier": cap,
                "capacity_tier_num": config.CAPACITY_TIERS.get(cap, 0),
                "db_type": r.get("db_type", ""),
                "release_patch": r.get("release_patch", ""),
                "release_family": r.get("release_family_version", ""),
            }
    console.print(f"  Loaded [cyan]{len(data):,}[/cyan] instances from prod_transactions CSV")
    return data


# ── Full scoring (with audit data) ──────────────────────────────────────────

def _compute_product_area_sums(
    table_counts: Dict[str, Optional[int]],
) -> Dict[str, Optional[int]]:
    """Compute product area aggregate row counts + active area count."""
    result = {}
    active = 0
    for area, tables in config.PRODUCT_AREAS.items():
        total = 0
        found = False
        for t in tables:
            v = table_counts.get(t)
            if v and v > 0:
                total += v
                found = True
        result[area] = total if found else None
        if found and total > 0:
            active += 1
    result["active"] = active
    return result


def _normalize(value, dim_min, dim_max):
    if not value or value <= 0:
        return 0.0
    if dim_max == dim_min:
        return 0.0
    return (value - dim_min) / (dim_max - dim_min)


def rank_instances_full(
    csv_data: Dict[str, dict],
    top_n: int = config.DEFAULT_TOP_N,
) -> List[Tuple[str, float, int, dict]]:
    """
    Full 18-dimension composite scoring.
    Returns [(instance, score, coverage, dims_dict), ...] sorted desc.
    """
    console.print("  Loading audit data files...")

    total_rows = _load_audit_xlsx(config.AUDIT_TOTAL_ROWS_XLSX, _parse_row_count)
    attachments = _load_audit_xlsx(config.AUDIT_ATTACHMENTS_XLSX, _parse_row_count)
    attach_docs = _load_audit_xlsx(config.AUDIT_ATTACH_DOCS_XLSX, _parse_row_count)
    flows_hm = _load_audit_xlsx(config.AUDIT_FLOWS_XLSX, _parse_flow_count)
    ecc_agents = _load_audit_xlsx(config.AUDIT_ECC_AGENTS_XLSX, _parse_row_count)
    user_prefs = _load_audit_xlsx(config.AUDIT_USER_PREFS_XLSX, _parse_row_count)
    table_data, _ = _load_table_stats(config.AUDIT_TABLE_STATS_XLSX)

    audit_sources = {
        "total_rows": total_rows, "attachments": attachments,
        "attach_docs": attach_docs, "flows": flows_hm,
        "ecc_agents": ecc_agents, "user_prefs": user_prefs,
    }
    audit_avail = sum(1 for v in audit_sources.values() if len(v) > 0)
    console.print(f"  Audit sources loaded: [cyan]{audit_avail}/6[/cyan] available")

    if audit_avail == 0 and len(table_data) == 0:
        console.print("  [yellow]No audit data found — falling back to CSV-only scoring[/yellow]")
        return rank_instances_csv_only(csv_data, top_n)

    # Build product-area counts
    pa = {i: _compute_product_area_sums(tc) for i, tc in table_data.items()}

    all_inst = set(csv_data) | set(table_data) | set(total_rows)

    # Assemble dimensions per instance
    dims_data = {}
    for inst in all_inst:
        p = pa.get(inst, {})
        td = csv_data.get(inst, {})
        d = {
            "total_rows": total_rows.get(inst),
            "txn_90d": td.get("txn_90d") if td.get("txn_90d", 0) > 0 else None,
            "db_gb": td.get("db_gb") if td.get("db_gb", 0) > 0 else None,
            "attach": attachments.get(inst),
            "attach_doc": attach_docs.get(inst),
            "itsm": p.get("itsm"),
            "cmdb": p.get("cmdb"),
            "itom": p.get("itom"),
            "events": p.get("events"),
            "hr": p.get("hr"),
            "secops": p.get("secops"),
            "sam": p.get("sam"),
            "flows": flows_hm.get(inst),
            "mids": ecc_agents.get(inst),
            "users": (
                p.get("users")
                if (p.get("users") or 0) > 0
                else table_data.get(inst, {}).get("sys user")
            ),
            "ecc_q": p.get("ecc_q"),
            "active": p.get("active"),
            "kb": p.get("kb"),
        }
        for k in d:
            if d[k] is not None and isinstance(d[k], (int, float)) and d[k] <= 0:
                d[k] = None
        dims_data[inst] = d

    # Compute normalization ranges
    W = config.SCORING_WEIGHTS
    ranges = {}
    for dim in W:
        vals = [
            dims_data[i][dim]
            for i in dims_data
            if dims_data[i].get(dim) and dims_data[i][dim] > 0
        ]
        ranges[dim] = (min(vals), max(vals)) if vals else (0, 1)

    # Score each instance
    scored = []
    for inst, d in dims_data.items():
        coverage = sum(1 for dim in W if d.get(dim) and d[dim] > 0)
        score = sum(
            W[dim] * _normalize(d[dim], *ranges[dim]) for dim in W
        )
        scored.append((inst, score, coverage, d))

    scored.sort(key=lambda x: x[1], reverse=True)
    return scored[:top_n] if top_n else scored


# ── CSV-only scoring (fallback) ──────────────────────────────────────────────

def rank_instances_csv_only(
    csv_data: Dict[str, dict],
    top_n: int = config.DEFAULT_TOP_N,
) -> List[Tuple[str, float, int, dict]]:
    """
    Simplified scoring using only prod_transactions CSV data.
    4 dimensions: txn_90d, db_gb, primary_table_gb, capacity_tier.
    """
    W = config.CSV_ONLY_WEIGHTS

    # Compute ranges
    ranges = {}
    for dim in W:
        if dim == "capacity_tier":
            vals = [d.get("capacity_tier_num", 0) for d in csv_data.values() if d.get("capacity_tier_num", 0) > 0]
        else:
            vals = [d.get(dim, 0) for d in csv_data.values() if d.get(dim, 0) > 0]
        ranges[dim] = (min(vals), max(vals)) if vals else (0, 1)

    scored = []
    for inst, d in csv_data.items():
        dims = {
            "txn_90d": d.get("txn_90d", 0),
            "db_gb": d.get("db_gb", 0),
            "primary_table_gb": d.get("primary_table_gb", 0),
            "capacity_tier": d.get("capacity_tier_num", 0),
        }
        coverage = sum(1 for v in dims.values() if v and v > 0)
        score = sum(
            W[dim] * _normalize(dims[dim], *ranges[dim]) for dim in W
        )
        # Store CSV dims in the same dict format
        full_dims = {
            "total_rows": None,
            "txn_90d": d.get("txn_90d"),
            "db_gb": d.get("db_gb"),
            "primary_table_gb": d.get("primary_table_gb"),
            "capacity_tier": d.get("capacity_tier"),
        }
        scored.append((inst, score, coverage, full_dims))

    scored.sort(key=lambda x: x[1], reverse=True)
    return scored[:top_n] if top_n else scored


# ── Public API ───────────────────────────────────────────────────────────────

def get_ranked_instances(
    top_n: int = config.DEFAULT_TOP_N,
    csv_path: str = None,
    csv_only: bool = False,
) -> List[Tuple[str, float, int, dict]]:
    """
    Main entry point: load data, score, and return ranked instances.
    Returns: [(instance_name, score, coverage, dimensions_dict), ...]
    """
    console.print("\n[bold cyan]Step 1: Ranking instances by complexity[/bold cyan]")
    csv_data = load_prod_transactions(csv_path)

    if not csv_data:
        console.print("[red]No CSV data loaded — cannot rank instances[/red]")
        return []

    if csv_only:
        console.print("  Using CSV-only scoring (4 dimensions)")
        ranked = rank_instances_csv_only(csv_data, top_n)
    else:
        ranked = rank_instances_full(csv_data, top_n)

    # Merge Prod Customers xlsx instances not already ranked
    prod_customers = load_prod_customers()
    ranked_names = {r[0] for r in ranked}
    extra_count = 0
    for inst, info in prod_customers.items():
        if inst not in ranked_names:
            # Give unscored instances score=0, coverage=0, empty dims
            ranked.append((inst, 0.0, 0, {}))
            extra_count += 1
    if extra_count:
        console.print(f"  Added [cyan]{extra_count:,}[/cyan] unscored instances from Prod Customers xlsx")

    if ranked:
        scored_only = [r for r in ranked if r[1] > 0]
        if scored_only:
            console.print(
                f"  [green]Total {len(ranked):,} instances[/green] "
                f"({len(scored_only):,} scored, range: {scored_only[0][1]:.4f} → {scored_only[-1][1]:.4f})"
            )
        else:
            console.print(f"  [green]Total {len(ranked):,} instances[/green] (no scores available)")
    return ranked
