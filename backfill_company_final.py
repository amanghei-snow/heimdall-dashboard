"""Backfill instance.company from cmdb_ci_service xlsx files using batched CASE updates."""
import os
import time
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

XLSX_PATHS = [
    "/Users/aman.ghei/Downloads/cmdb_ci_service.xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (1).xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (2).xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (3).xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (4).xlsx",
    "/Users/aman.ghei/CascadeProjects/ruckus-aggregator/data/cmdb_ci_service (2).xlsx",
]


def build_mapping():
    mapping = {}
    for path in XLSX_PATHS:
        if not os.path.exists(path):
            continue
        print(f"Reading {os.path.basename(path)} ...")
        wb = load_workbook(path, data_only=False)
        ws = wb.active
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {v: i for i, v in enumerate(header) if v}
        name_col = headers.get("Name")
        company_col = headers.get("Company")
        suffix_col = headers.get("Suffix (calc)")
        if name_col is None:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[name_col]
            company = row[company_col] if company_col is not None else None
            suffix = row[suffix_col] if suffix_col is not None else None
            if not name:
                continue
            inst = suffix
            if not inst and isinstance(name, str):
                if name.lower().startswith("snc instance - "):
                    inst = name[len("SNC Instance - "):]
                else:
                    inst = name
            if not inst:
                continue
            inst = str(inst).strip().lower()
            if company and str(company).strip() and inst not in mapping:
                mapping[inst] = str(company).strip()
    return mapping


def update_blank_companies(engine_url, mapping, batch_size=500):
    label = engine_url.split("@")[-1].split("/")[0]
    print(f"\nUpdating {label} ...")
    e = create_engine(engine_url, echo=False, pool_pre_ping=True)
    with e.connect() as c:
        res = c.execute(text(
            "SELECT instance FROM instances WHERE company IS NULL OR company = '' OR company = 'Unknown'"
        )).fetchall()
        blanks = {r[0] for r in res}
        print(f"  {len(blanks)} blank companies")

        to_update = {k: v for k, v in mapping.items() if k in blanks}
        print(f"  {len(to_update)} matched by xlsx")
        if not to_update:
            return 0

        items = list(to_update.items())
        total_updated = 0
        for i in range(0, len(items), batch_size):
            batch = items[i:i + batch_size]
            # Build one UPDATE ... CASE statement with bind params
            whens = "\n".join([f"    WHEN :i{idx} THEN :c{idx}" for idx in range(len(batch))])
            in_list = ", ".join([f":i{idx}" for idx in range(len(batch))])
            stmt = text(
                f"UPDATE instances\n"
                f"SET company = CASE instance\n{whens}\n    ELSE company\nEND\n"
                f"WHERE instance IN ({in_list})\n"
                f"  AND (company IS NULL OR company = '' OR company = 'Unknown')"
            )
            params = {}
            for idx, (inst, co) in enumerate(batch):
                params[f"i{idx}"] = inst
                params[f"c{idx}"] = co
            result = c.execute(stmt, params)
            c.commit()
            total_updated += result.rowcount
            print(f"  updated {total_updated} / {len(items)}")
    return total_updated


if __name__ == "__main__":
    t0 = time.time()
    mapping = build_mapping()
    print(f"\nBuilt {len(mapping)} unique instance->company mappings in {time.time()-t0:.2f}s")

    local_url = os.getenv("DATABASE_URL", "mysql+pymysql://ruckus:ruckus@localhost:3306/ruckus_aggregator")
    render_url = os.getenv(
        "RENDER_DATABASE_URL",
        "postgresql://heimdall:F8PKD2ASPB2PDsgvgfa9oplTXUGgLmW4@dpg-dad12p710e5c73cs6qgg-a.oregon-postgres.render.com/heimdall_p8ps",
    )

    update_blank_companies(local_url, mapping)
    update_blank_companies(render_url, mapping)
    print(f"\nFinished in {time.time()-t0:.2f}s")
