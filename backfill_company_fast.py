"""Fast backfill of instance.company from cmdb_ci_service xlsx files."""
import os
import json
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

XLSX_PATHS = [
    "/Users/aman.ghei/Downloads/cmdb_ci_service.xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (1).xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (2).xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (3).xlsx",
    "/Users/aman.ghei/Downloads/cmdb_ci_service (4).xlsx",
]


def build_mapping():
    mapping = {}
    for path in XLSX_PATHS:
        if not os.path.exists(path):
            continue
        print(f"Reading {path} ...")
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        # Read header
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {val: idx for idx, val in enumerate(header_row) if val}
        name_col = headers.get("Name") or headers.get("name")
        company_col = headers.get("Company") or headers.get("company")
        suffix_col = headers.get("Suffix (calc)") or headers.get("Suffix")
        if name_col is None:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[name_col]
            company = row[company_col] if company_col is not None else None
            suffix = row[suffix_col] if suffix_col is not None else None
            if not name:
                continue
            inst = suffix
            if not inst and isinstance(name, str):
                if name.lower().startswith("snc instance - "):
                    inst = name[len("SNC Instance - "):]
                else:
                    inst = name
            if not inst:
                continue
            inst = str(inst).strip().lower()
            if company and str(company).strip() and inst not in mapping:
                mapping[inst] = str(company).strip()
    return mapping


def update_blank_companies(engine_url, mapping):
    print(f"\nConnecting {engine_url.split('@')[-1]} ...")
    e = create_engine(engine_url, echo=False)
    with e.connect() as c:
        # find blank instances
        res = c.execute(text(
            "SELECT instance FROM instances WHERE company IS NULL OR company = '' OR company = 'Unknown'"
        )).fetchall()
        blanks = {r[0] for r in res}
        print(f"  {len(blanks)} blank companies")
        to_update = {k: v for k, v in mapping.items() if k in blanks}
        print(f"  {len(to_update)} matches from xlsx")
        if not to_update:
            return 0

        # batch update with executemany-style
        items = list(to_update.items())
        batch_size = 1000
        updated = 0
        for i in range(0, len(items), batch_size):
            batch = items[i:i + batch_size]
            stmt = text(
                "UPDATE instances SET company = :company WHERE instance = :instance AND (company IS NULL OR company = '' OR company = 'Unknown')"
            )
            c.execute(stmt, [{"company": co, "instance": inst} for inst, co in batch])
            c.commit()
            updated += len(batch)
            print(f"  updated {updated} / {len(items)}")
    return len(to_update)


if __name__ == "__main__":
    mapping = build_mapping()
    print(f"\nBuilt mapping with {len(mapping)} entries")

    local_url = os.getenv("DATABASE_URL", "mysql+pymysql://ruckus:ruckus@localhost:3306/ruckus_aggregator")
    render_url = os.getenv(
        "RENDER_DATABASE_URL",
        "postgresql://heimdall:F8PKD2ASPB2PDsgvgfa9oplTXUGgLmW4@dpg-dad12p710e5c73cs6qgg-a.oregon-postgres.render.com/heimdall_p8ps",
    )

    update_blank_companies(local_url, mapping)
    update_blank_companies(render_url, mapping)
    print("Done")
